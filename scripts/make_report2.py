#!/usr/bin/env python
"""「合集 / 专辑」TMDB × 115 对照表（**只读**，不碰网盘）。

输入:
    data/state/container_scan.json      scripts/scan_containers.py
    data/state/title_verify.json        scripts/verify_titles.py
    data/state/115_cids.json            scripts/fetch_115_cids.py
    data/state/tmdb_collections.json    scripts/fetch_tmdb_collections.py
    data/state/inbox.json               scripts/fetch_inbox.py

输出: ~/WorkBuddy/aix8-pan/合集_专辑_TMDB核对表.xlsx
    Sheet1 合集      逐作品一行 —— TMDB 合集 + 网盘系列 + TMDB 单片 + 网盘单片
                     **已上映但网盘缺的，也补一行**，问题说明标「暂无网盘资源」
    Sheet2 专辑      逐作品一行 —— 仅单片维度
    Sheet3 系列汇总  逐系列一行 —— 用「已上映数 vs 网盘数」标出还缺哪些

与 v2 的差异（2026-10-07 修订）:
    1. 新增 4 列伴随文件清单：movie_pan_nfo / _poster / _fanart / _clearlogo
    2. movie_pan_size → movie_pan_size(GB)，单位由 MB 改 G（两位小数）
    3. 合集 sheet 补「已上映·网盘暂无」行，问题说明 = 暂无网盘资源
    4. movie_pan_tags 去掉分辨率段（分辨率已有独立列 movie_pan_res，避免两列重复）
"""
from __future__ import annotations

import json
from collections import defaultdict
from itertools import groupby
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

import sys
sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
from aix8pan.config import state_path

from aix8pan.core.naming_spec import TECH_ORDER, parse_artwork
from aix8pan.core.parser import extract_tech_fields

SCAN = state_path("container_scan.json")
TVF = state_path("title_verify.json")
LKF = state_path("115_cids.json")
TCF = state_path("tmdb_collections.json")
IBF = state_path("inbox.json")
OUT = Path.home() / "WorkBuddy/aix8-pan/合集_专辑_TMDB核对表.xlsx"

C_HEAD = "1F3864"
C_HEAD_TXT = "FFFFFF"
C_OK = "E2EFDA"
C_OK_TXT = "375623"
C_BAD = "FCE4E4"
C_BAD_TXT = "9C0006"
C_MID = "FFF2CC"
C_MID_TXT = "7F6000"
C_LINK = "0563C1"
C_NA = "F2F2F2"
C_NA_TXT = "808080"
C_GRP = "D9E1F2"

THIN = Side(style="thin", color="D0D7E5")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(vertical="center", wrap_text=True)
CENTER = Alignment(horizontal="center", vertical="center")

MISSING_NOTE = "暂无网盘资源"


def style_header(ws, ncol: int, row: int = 1, height: int = 34, freeze_col: int = 4):
    ws.row_dimensions[row].height = height
    for c in range(1, ncol + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = PatternFill("solid", fgColor=C_HEAD)
        cell.font = Font(bold=True, color=C_HEAD_TXT, size=10)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER
    ws.freeze_panes = f"{get_column_letter(freeze_col)}{row + 1}"


def put_link(ws, row: int, col: int, url: str, text: str | None = None):
    cell = ws.cell(row=row, column=col)
    if text is not None:
        cell.value = text
    if url:
        cell.hyperlink = url
        cell.font = Font(size=10, color=C_LINK, underline="single")


def verdict(r: dict) -> tuple[str, str, str]:
    if r["conform"]:
        return "✅ 符合", C_OK, C_OK_TXT
    if r["n_err"] == 0:
        return "⚠️ 基本符合", C_MID, C_MID_TXT
    return "❌ 不符合", C_BAD, C_BAD_TXT


def join_issues(r: dict) -> str:
    parts = []
    for e in r["errors"]:
        parts.append("【必须整改】" + e)
    for w in r["warns"]:
        parts.append("【建议整改】" + w)
    for i in r["infos"]:
        parts.append("【提示】" + i)
    return "\n".join(parts)


def tech_no_res(tech_str: str) -> str:
    """技术规格串 → 去掉分辨率段（分辨率已单列）。

    '1080p AC3 / 2160p DTSHD-MA' → 'AC3 / DTSHD-MA'
    """
    segs: list[str] = []
    for seg in str(tech_str or "").split(" / "):
        f = extract_tech_fields(seg)
        s = " ".join(f[k] for k in TECH_ORDER if k != "resolution" and f.get(k))
        if s and s not in segs:
            segs.append(s)
    return " / ".join(segs)


def companion_flags(r: dict) -> dict[str, bool]:
    """该作品目录里 nfo / 海报 / 背景 / Logo 是否存在（任一版本有即算有）。"""
    kinds = {parse_artwork(n) for n in r.get("images") or []}
    return {
        "nfo": bool(r.get("nfos")),
        "poster": "poster" in kinds,
        "fanart": "fanart" in kinds,
        "clearlogo": "clearlogo" in kinds,
    }


def subs_label(r: dict) -> str:
    """字幕列：简中/英文 覆盖情况。"""
    subs = [str(s).lower() for s in (r.get("subs") or [])]
    if not subs:
        return "❌ 无"
    zh = any((".zh" in s) or ("chs" in s) or ("简" in s) for s in subs)
    en = any((".en." in s) or s.endswith(".en.srt") or ("eng" in s) for s in subs)
    if zh and en:
        return "简中+英文"
    if zh:
        return "仅简中"
    if en:
        return "仅英文"
    return f"其他({len(subs)})"


# 专辑 sheet 的列（合集 sheet 的后半段与之同名同序）
MOVIE_HEADS = [
    "tmdb_movie_id", "tmdb_movie_name", "tmdb_movie_year", "tmdb_movie_release",
    "tmdb_movie_link",
    "movie_pan_name", "movie_pan_link", "movie_pan_main", "movie_pan_size(GB)",
    "movie_pan_res",
    "movie_pan_tags", "movie_pan_count", "movie_pan_versions", "movie_pan_media_count",
    "movie_pan_nfo", "movie_pan_poster", "movie_pan_fanart", "movie_pan_clearlogo",
    "movie_pan_subs",
    "movie_pan_folder_ok", "movie_pan_file_ok", "movie_pan_art_ok",
    "是否符合规范", "问题说明",
]

COLL_HEADS = [
    "tmdb_collection_id", "tmdb_collection_name", "tmdb_collection_link",
    "tmdb_collection_total", "tmdb_collection_released",
    "collection_pan_link", "collection_pan_count",
    "collection_missing_list",
]

TICK = {True: "✅", False: "❌"}

# 补缺行里「网盘侧」的全部列（整块置灰）
PAN_COLS = (
    "movie_pan_name", "movie_pan_link", "movie_pan_main", "movie_pan_size(GB)",
    "movie_pan_res",
    "movie_pan_tags", "movie_pan_count", "movie_pan_versions",
    "movie_pan_media_count", "movie_pan_nfo", "movie_pan_poster",
    "movie_pan_fanart", "movie_pan_clearlogo", "movie_pan_subs",
    "movie_pan_folder_ok", "movie_pan_file_ok", "movie_pan_art_ok",
)


def main() -> int:
    data = json.loads(SCAN.read_text(encoding="utf-8"))
    rows = data["rows"]

    tv = {}
    if TVF.exists():
        for r in json.loads(TVF.read_text(encoding="utf-8")):
            tv[(r["container"], r["folder"])] = r

    links = {}
    if LKF.exists():
        links = {k: v["url"] for k, v in
                 json.loads(LKF.read_text(encoding="utf-8"))["map"].items()}

    tc = json.loads(TCF.read_text(encoding="utf-8")) if TCF.exists() else {}
    by_movie = tc.get("by_movie", {})
    series_info = tc.get("series", {})
    colls = tc.get("collections", {})
    today = tc.get("today", "")

    # 「0-待整理」区已有（用于区分「没资源」和「没归位」）
    inbox: dict[str, str] = {}
    if IBF.exists():
        for it in json.loads(IBF.read_text(encoding="utf-8")).get("items", []):
            if it.get("tmdb_id"):
                inbox[str(it["tmdb_id"])] = it.get("title") or it.get("name") or "该片"

    # 每个系列在网盘里的作品 id 集合
    pan_ids: dict[str, set[str]] = defaultdict(set)
    for r in rows:
        pan_ids[r["container"]].add(str(r["tmdb_id"]))

    # 单片上映日 / 名称（从各合集的 parts 汇总，作为 TMDB 侧的第一手信息）
    part_date: dict[str, str] = {}
    part_title: dict[str, str] = {}
    for c in colls.values():
        for p in c["parts"]:
            part_date.setdefault(p["id"], p["release_date"])
            part_title.setdefault(p["id"], p["title"])

    def coll_block(container: str) -> dict:
        """给一个系列算出 collection_* 各项。"""
        cs = series_info.get(container) or {}
        cid = str(cs.get("collection_id") or "")
        c = colls.get(cid) or {}
        if not cid or not c:
            return {"id": "—", "name": "—（TMDB 无对应合集）", "url": "", "text": "—",
                    "count": "—", "released": "—", "unreleased": "—",
                    "pan_count": len(pan_ids.get(container, ())),
                    "missing": "—", "missing_list": "", "miss_items": []}
        have = pan_ids.get(container, set())
        miss = [p for p in c["parts"] if p["released"] and p["id"] not in have]
        miss_txt = "\n".join(f"{p['title']} ({p['release_date'][:4]})" for p in miss)
        return {
            "id": cid,
            "name": c["name_zh"] or c["original_name"] or "—",
            "url": c["url"],
            "text": c["name_zh"] or c["original_name"] or "—",
            "count": c["count_total"],
            "released": c["count_released"],
            "unreleased": c["count_unreleased"],
            "pan_count": len(have),
            "missing": len(miss),
            "missing_list": miss_txt or "—",
            "miss_items": miss,
        }

    def movie_cells(r: dict) -> list:
        """网盘里**已有**的作品行（后半段列）。"""
        mid = str(r["tmdb_id"])
        bm = by_movie.get(mid) or {}
        rd = bm.get("release_date") or part_date.get(mid) or ""
        vt, _, _ = verdict(r)
        cf = companion_flags(r)
        return [
            mid,
            bm.get("title") or part_title.get(mid) or r["title"],
            (rd[:4] if rd else r["year"]) or "—",
            rd or "—",
            f"https://www.themoviedb.org/movie/{mid}",
            r["folder_raw"],
            links.get(r["path"], ""),
            r.get("main_file") or "—",
            round((r.get("size_mb") or 0) / 1024.0, 2),
            r["resolution"] or "—",
            tech_no_res(r.get("tech_all") or "") or "—",
            r.get("files_total", "—"),
            r["version_count"],
            r["media_count"],
            TICK[cf["nfo"]], TICK[cf["poster"]], TICK[cf["fanart"]], TICK[cf["clearlogo"]],
            subs_label(r),
            TICK[r["folder_ok"]], TICK[r["file_ok"]], TICK[r["art_ok"]],
            vt,
            join_issues(r) or "—",
        ]

    def missing_movie_cells(m: dict) -> list:
        """已上映但网盘缺的 TMDB 单片 —— 补一行，问题说明标「暂无网盘资源」。"""
        mid = str(m["id"])
        rd = m.get("release_date") or ""
        bm = by_movie.get(mid) or {}
        note = MISSING_NOTE
        if mid in inbox:
            note = f"{MISSING_NOTE}（0-待整理已有《{inbox[mid]}》，待归位）"
        return [
            mid,
            m.get("title") or bm.get("title") or part_title.get(mid) or "—",
            (rd[:4] if rd else "—"),
            rd or "—",
            f"https://www.themoviedb.org/movie/{mid}",
            "—", "", "—", "—", "—", "—", "—", "—", "—",
            "—", "—", "—", "—",
            "—",
            "—", "—", "—",
            "⬜ 暂无资源",
            note,
        ]

    wb = Workbook()

    # ══════════════════ Sheet1 合集 ══════════════════
    ws = wb.active
    ws.title = "合集"
    heads = ["序号"] + COLL_HEADS + MOVIE_HEADS
    ws.append(heads)
    style_header(ws, len(heads), freeze_col=1)

    coll_rows = sorted([r for r in rows if r["category"] == "合集"],
                       key=lambda r: (r["container"], r["year"] or ""))
    coll_cache: dict[str, dict] = {}
    idx = 0
    n_missing_added = 0
    miss_log: list[str] = []

    def _paint_row(n: int):
        for c in range(1, len(heads) + 1):
            cell = ws.cell(row=n, column=c)
            cell.border = BORDER
            cell.alignment = WRAP
            cell.font = Font(size=10)

    def _center(names, row_n):
        for name in names:
            ws.cell(row=row_n, column=heads.index(name) + 1).alignment = CENTER

    for cont, grp in groupby(coll_rows, key=lambda r: r["container"]):
        if cont not in coll_cache:
            coll_cache[cont] = coll_block(cont)
        cb = coll_cache[cont]
        grp = list(grp)

        for r in grp:
            idx += 1
            ws.append([idx,
                       cb["id"], cb["name"], cb["url"],
                       cb["count"], cb["released"],
                       links.get(f"/115/01-电影/{cont}", ""), cb["pan_count"],
                       cb["missing_list"]]
                      + movie_cells(r))
            n = ws.max_row
            _paint_row(n)
            put_link(ws, n, heads.index("tmdb_collection_link") + 1, cb["url"])
            put_link(ws, n, heads.index("collection_pan_link") + 1,
                     links.get(f"/115/01-电影/{cont}", ""))
            put_link(ws, n, heads.index("tmdb_movie_link") + 1,
                     f"https://www.themoviedb.org/movie/{r['tmdb_id']}")
            put_link(ws, n, heads.index("movie_pan_link") + 1, links.get(r["path"], ""))
            vcell = ws.cell(row=n, column=heads.index("是否符合规范") + 1)
            vt, fill, txt = verdict(r)
            vcell.fill = PatternFill("solid", fgColor=fill)
            vcell.font = Font(size=10, bold=True, color=txt)
            vcell.alignment = CENTER
            for name in ("movie_pan_folder_ok", "movie_pan_file_ok", "movie_pan_art_ok"):
                cc = ws.cell(row=n, column=heads.index(name) + 1)
                cc.alignment = CENTER
                cc.font = Font(size=11, color=C_OK_TXT if cc.value == "✅" else C_BAD_TXT)
            for name in ("movie_pan_nfo", "movie_pan_poster", "movie_pan_fanart",
                         "movie_pan_clearlogo"):
                cc = ws.cell(row=n, column=heads.index(name) + 1)
                cc.alignment = CENTER
                cc.font = Font(size=11, color=C_OK_TXT if cc.value == "✅" else C_BAD_TXT)
            sc = ws.cell(row=n, column=heads.index("movie_pan_subs") + 1)
            sc.alignment = CENTER
            sc.font = Font(size=10, bold=(sc.value == "❌ 无"),
                           color=C_BAD_TXT if sc.value == "❌ 无" else "334155")
            _center(("序号", "tmdb_collection_id", "tmdb_collection_total",
                     "tmdb_collection_released",
                     "collection_pan_count", "tmdb_movie_id", "tmdb_movie_year",
                     "movie_pan_size(GB)", "movie_pan_res", "movie_pan_count",
                     "movie_pan_versions", "movie_pan_media_count"), n)
            for name in ("tmdb_movie_id", "tmdb_movie_year", "tmdb_collection_id"):
                ws.cell(row=n, column=heads.index(name) + 1).number_format = "@"
            ws.cell(row=n, column=heads.index("movie_pan_size(GB)") + 1).number_format = "0.00"

        # ── 已上映但网盘暂无的，补行 ──
        for m in cb["miss_items"]:
            idx += 1
            n_missing_added += 1
            ws.append([idx,
                       cb["id"], cb["name"], cb["url"],
                       cb["count"], cb["released"],
                       links.get(f"/115/01-电影/{cont}", ""), cb["pan_count"],
                       cb["missing_list"]]
                      + missing_movie_cells(m))
            n = ws.max_row
            _paint_row(n)
            put_link(ws, n, heads.index("tmdb_collection_link") + 1, cb["url"])
            put_link(ws, n, heads.index("collection_pan_link") + 1,
                     links.get(f"/115/01-电影/{cont}", ""))
            put_link(ws, n, heads.index("tmdb_movie_link") + 1,
                     f"https://www.themoviedb.org/movie/{m['id']}")
            # 灰底：网盘侧全空
            for name in PAN_COLS:
                ws.cell(row=n, column=heads.index(name) + 1).fill = \
                    PatternFill("solid", fgColor=C_NA)
                ws.cell(row=n, column=heads.index(name) + 1).font = \
                    Font(size=10, color=C_NA_TXT)
            vcell = ws.cell(row=n, column=heads.index("是否符合规范") + 1)
            vcell.value = "⬜ 暂无资源"
            vcell.fill = PatternFill("solid", fgColor=C_NA)
            vcell.font = Font(size=10, bold=True, color=C_NA_TXT)
            vcell.alignment = CENTER
            ncell = ws.cell(row=n, column=heads.index("问题说明") + 1)
            ncell.fill = PatternFill("solid", fgColor=C_MID)
            ncell.font = Font(size=10, bold=True, color=C_MID_TXT)
            _center(("序号", "tmdb_collection_id", "tmdb_collection_total",
                     "tmdb_collection_released",
                     "collection_pan_count", "tmdb_movie_id", "tmdb_movie_year"), n)
            for name in ("tmdb_movie_id", "tmdb_movie_year", "tmdb_collection_id"):
                ws.cell(row=n, column=heads.index(name) + 1).number_format = "@"
            miss_log.append(f"{cont.replace('合集/', '', 1)} → {m['title']} ({m['release_date'][:4]})")

    widths = {
        "序号": 5,
        "tmdb_collection_id": 11, "tmdb_collection_name": 20, "tmdb_collection_link": 44,
        "tmdb_collection_total": 10, "tmdb_collection_released": 11,
        "collection_pan_link": 44,
        "collection_pan_count": 12,
        "collection_missing_list": 34,
        "tmdb_movie_id": 11, "tmdb_movie_name": 22, "tmdb_movie_year": 7,
        "tmdb_movie_release": 13, "tmdb_movie_link": 40,
        "movie_pan_name": 34, "movie_pan_link": 40, "movie_pan_main": 46,
        "movie_pan_size(GB)": 13,
        "movie_pan_res": 12, "movie_pan_tags": 22, "movie_pan_count": 10,
        "movie_pan_versions": 8, "movie_pan_media_count": 9,
        "movie_pan_nfo": 9, "movie_pan_poster": 9, "movie_pan_fanart": 9,
        "movie_pan_clearlogo": 10, "movie_pan_subs": 11,
        "movie_pan_folder_ok": 8, "movie_pan_file_ok": 8, "movie_pan_art_ok": 8,
        "是否符合规范": 12, "问题说明": 40,
    }
    for idx2, h in enumerate(heads, 1):
        ws.column_dimensions[get_column_letter(idx2)].width = widths.get(h, 14)
    ws.auto_filter.ref = f"A1:{get_column_letter(len(heads))}{ws.max_row}"

    # ══════════════════ Sheet2 专辑 ══════════════════
    ws2 = wb.create_sheet("专辑")
    h2 = ["序号"] + MOVIE_HEADS
    ws2.append(h2)
    style_header(ws2, len(h2), freeze_col=1)
    alb = sorted([r for r in rows if r["category"] == "专辑"], key=lambda r: r["title"])
    for i, r in enumerate(alb, 1):
        ws2.append([i] + movie_cells(r))
        n = ws2.max_row
        for c in range(1, len(h2) + 1):
            cell = ws2.cell(row=n, column=c)
            cell.border = BORDER
            cell.alignment = WRAP
            cell.font = Font(size=10)
        put_link(ws2, n, h2.index("tmdb_movie_link") + 1,
                 f"https://www.themoviedb.org/movie/{r['tmdb_id']}")
        put_link(ws2, n, h2.index("movie_pan_link") + 1, links.get(r["path"], ""))
        vcell = ws2.cell(row=n, column=h2.index("是否符合规范") + 1)
        vt, fill, txt = verdict(r)
        vcell.fill = PatternFill("solid", fgColor=fill)
        vcell.font = Font(size=10, bold=True, color=txt)
        vcell.alignment = CENTER
        for name in ("movie_pan_folder_ok", "movie_pan_file_ok", "movie_pan_art_ok",
                     "movie_pan_nfo", "movie_pan_poster", "movie_pan_fanart",
                     "movie_pan_clearlogo"):
            cc = ws2.cell(row=n, column=h2.index(name) + 1)
            cc.alignment = CENTER
            cc.font = Font(size=11, color=C_OK_TXT if cc.value == "✅" else C_BAD_TXT)
        sc2 = ws2.cell(row=n, column=h2.index("movie_pan_subs") + 1)
        sc2.alignment = CENTER
        sc2.font = Font(size=10, bold=(sc2.value == "❌ 无"),
                        color=C_BAD_TXT if sc2.value == "❌ 无" else "334155")
        for name in ("序号", "tmdb_movie_id", "tmdb_movie_year",
                     "movie_pan_size(GB)", "movie_pan_res", "movie_pan_count",
                     "movie_pan_versions", "movie_pan_media_count"):
            ws2.cell(row=n, column=h2.index(name) + 1).alignment = CENTER
        for name in ("tmdb_movie_id", "tmdb_movie_year"):
            ws2.cell(row=n, column=h2.index(name) + 1).number_format = "@"
        ws2.cell(row=n, column=h2.index("movie_pan_size(GB)") + 1).number_format = "0.00"
    for idx2, h in enumerate(h2, 1):
        ws2.column_dimensions[get_column_letter(idx2)].width = widths.get(h, 14)
    ws2.auto_filter.ref = f"A1:{get_column_letter(len(h2))}{ws2.max_row}"

    # ══════════════════ Sheet3 系列汇总 ══════════════════
    ws3 = wb.create_sheet("系列汇总")
    h3 = ["序号", "网盘系列目录", "collection_pan_link", "tmdb_collection_id",
          "tmdb_collection_name", "tmdb_collection_link", "tmdb_collection_total",
          "tmdb_collection_released",
          "collection_pan_count", "collection_missing_list",
          "待整理区已有（待归位）", "整库缺失（TMDB 已上映但网盘无）",
          "网盘多出（不在 TMDB 合集内）", "异常提示"]
    ws3.append(h3)
    style_header(ws3, len(h3), height=44)
    k = 0
    miss_series = 0
    absent_series = 0
    all_missing: list[str] = []
    all_absent: list[str] = []
    for cont in sorted({r["container"] for r in rows if r["category"] == "合集"}):
        k += 1
        cb = coll_block(cont)
        have = pan_ids.get(cont, set())
        cid = str((series_info.get(cont) or {}).get("collection_id") or "")
        c = colls.get(cid) or {}
        extra = ""
        if c:
            member = {p["id"] for p in c["parts"]}
            extra_ids = have - member
            if extra_ids:
                titles = [f"{r['title']} ({r['year']})" for r in coll_rows
                          if r["container"] == cont and str(r["tmdb_id"]) in extra_ids]
                extra = "\n".join(titles)
        in_box = [p for p in cb["miss_items"] if p["id"] in inbox]
        absent = [p for p in cb["miss_items"] if p["id"] not in inbox]
        in_box_txt = "\n".join(f"{p['title']} ({p['release_date'][:4]})" for p in in_box)
        absent_txt = "\n".join(f"{p['title']} ({p['release_date'][:4]})" for p in absent)

        note = ""
        si = series_info.get(cont) or {}
        if si.get("mixed"):
            note = "⚠ 该系列下出现多个 TMDB 合集：" + "、".join(
                f"{v} 部归 {k}" for k, v in (si.get("votes") or {}).items())
        if not si.get("collection_id"):
            note = "⚠ TMDB 无对应合集，仅按网盘目录统计"
        ws3.append([k, cont.replace("合集/", "", 1),
                    links.get(f"/115/01-电影/{cont}", ""),
                    cb["id"], cb["name"], cb["url"],
                    cb["count"], cb["released"],
                    cb["pan_count"], cb["missing_list"],
                    in_box_txt or "—", absent_txt or "—", extra or "—",
                    note or "—"])
        n = ws3.max_row
        for cc in range(1, len(h3) + 1):
            cell = ws3.cell(row=n, column=cc)
            cell.border = BORDER
            cell.alignment = WRAP
            cell.font = Font(size=10)
        if in_box:
            ws3.cell(row=n, column=11).font = Font(size=10, color=C_MID_TXT)
        if absent:
            ws3.cell(row=n, column=12).font = Font(size=10, bold=True, color=C_BAD_TXT)
            absent_series += 1
            all_absent.append(f"{cont.replace('合集/', '', 1)} 缺 {len(absent)} 部")
        if note:
            ws3.cell(row=n, column=len(h3)).font = Font(size=10, color=C_MID_TXT)
        put_link(ws3, n, 3, links.get(f"/115/01-电影/{cont}", ""))
        put_link(ws3, n, 6, cb["url"])
        if isinstance(cb["missing"], int) and cb["missing"] > 0:
            miss_series += 1
            all_missing.append(f"{cont.replace('合集/', '', 1)} 缺 {cb['missing']} 部")
        if extra:
            ws3.cell(row=n, column=13).font = Font(size=10, color=C_MID_TXT)
        for name in ("序号", "tmdb_collection_id", "tmdb_collection_total",
                     "tmdb_collection_released",
                     "collection_pan_count"):
            ws3.cell(row=n, column=h3.index(name) + 1).alignment = CENTER
        ws3.cell(row=n, column=4).number_format = "@"
    for idx2, w in enumerate([5, 30, 42, 11, 20, 44, 10, 11, 12, 40, 30, 30, 30, 40], 1):
        ws3.column_dimensions[get_column_letter(idx2)].width = w
    ws3.auto_filter.ref = f"A1:{get_column_letter(len(h3))}{ws3.max_row}"

    wb.save(OUT)
    print(f"saved {OUT}")
    print(f"合集 {len(coll_rows)} 行 + 补缺 {n_missing_added} 行"
          f"｜专辑 {len(alb)} 行｜系列 {k} 个")
    print(f"与网盘系列有缺口的系列 {miss_series} 个" + (f"：{'；'.join(all_missing)}" if all_missing else ""))
    print(f"其中『整库确实没有』的系列 {absent_series} 个" + (f"：{'；'.join(all_absent)}" if all_absent else ""))
    for line in miss_log:
        print("  补行：" + line)
    if today:
        print(f"（‘已上映’基准日 {today}）")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
